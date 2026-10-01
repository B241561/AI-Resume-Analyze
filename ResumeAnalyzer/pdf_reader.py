from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document

from config import MAX_PDF_MB


class PdfReadError(Exception):
    """Raised when a resume file cannot be read safely."""


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_text_from_file(file_path: str | Path) -> str:
    path = Path(file_path)
    _validate_resume_file(path)

    if path.suffix.lower() == ".pdf":
        return extract_text_from_pdf(path)
    if path.suffix.lower() == ".docx":
        return extract_text_from_docx(path)

    raise PdfReadError("Please select a PDF or DOCX resume.")


def _validate_resume_file(path: Path) -> None:
    if not path.exists():
        raise PdfReadError("Selected resume file does not exist.")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise PdfReadError("Please select a PDF or DOCX resume.")
    if path.stat().st_size > MAX_PDF_MB * 1024 * 1024:
        raise PdfReadError(f"Resume file is too large. Maximum size is {MAX_PDF_MB} MB.")


def extract_text_from_pdf(pdf_path: str | Path) -> str:
    path = Path(pdf_path)
    _validate_resume_file(path)
    if path.suffix.lower() != ".pdf":
        raise PdfReadError("Please select a PDF resume.")

    try:
        document = fitz.open(path)
    except Exception as exc:
        raise PdfReadError("Could not open the PDF. It may be corrupted or password protected.") from exc

    pages: list[str] = []
    try:
        for page in document:
            pages.append(page.get_text("text"))
    finally:
        document.close()

    text = "\n".join(pages).strip()
    if len(text) < 50:
        raise PdfReadError("Could not extract enough text. Use a text-based resume PDF.")
    return text


def extract_text_from_docx(docx_path: str | Path) -> str:
    path = Path(docx_path)
    _validate_resume_file(path)
    if path.suffix.lower() != ".docx":
        raise PdfReadError("Please select a DOCX resume.")

    try:
        document = Document(path)
    except Exception as exc:
        raise PdfReadError("Could not open the DOCX file. It may be corrupted or unsupported.") from exc

    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]

    table_text: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_text.append(" | ".join(cells))

    text = "\n".join(paragraphs + table_text).strip()
    if len(text) < 50:
        raise PdfReadError("Could not extract enough text from the DOCX resume.")
    return text
